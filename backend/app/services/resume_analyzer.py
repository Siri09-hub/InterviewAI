import json
import logging
import os
import re
from pathlib import Path

from dotenv import load_dotenv
from google import genai
from google.genai import errors as genai_errors
from pypdf import PdfReader
from docx import Document

from app.services.resume_storage import resume_file_path


load_dotenv()

logger = logging.getLogger("interviewai.resume_analyzer")

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
GEMINI_MODEL = "gemini-2.5-flash"

MAX_RESUME_TEXT_LENGTH = 15000
MAX_ANALYSIS_ATTEMPTS = 3


class ResumeAnalysisError(Exception):
    """Raised when resume extraction or AI analysis fails."""

    pass


# --------------------------------------------------
# FILE TEXT EXTRACTION
# --------------------------------------------------

def _extract_pdf_text(
    path: Path,
) -> str:
    try:
        reader = PdfReader(str(path))

        pages: list[str] = []

        for page in reader.pages:
            text = page.extract_text() or ""

            if text.strip():
                pages.append(text)

        return "\n".join(pages)

    except Exception as exc:
        logger.exception(
            "PDF text extraction failed for %s",
            path,
        )

        raise ResumeAnalysisError(
            "Could not extract text from the PDF resume."
        ) from exc


def _extract_docx_text(
    path: Path,
) -> str:
    try:
        document = Document(str(path))

        paragraphs = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Also read table content because many resumes
        # store skills/education in tables.
        table_lines: list[str] = []

        for table in document.tables:
            for row in table.rows:
                values = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if values:
                    table_lines.append(
                        " | ".join(values)
                    )

        return "\n".join(
            paragraphs + table_lines
        )

    except Exception as exc:
        logger.exception(
            "DOCX text extraction failed for %s",
            path,
        )

        raise ResumeAnalysisError(
            "Could not extract text from the DOCX resume."
        ) from exc


def extract_resume_text(
    path: Path,
) -> str:
    if not path.exists():
        raise ResumeAnalysisError(
            "Stored resume file was not found."
        )

    extension = path.suffix.lower()

    if extension == ".pdf":
        text = _extract_pdf_text(path)

    elif extension == ".docx":
        text = _extract_docx_text(path)

    else:
        raise ResumeAnalysisError(
            "Only PDF and DOCX resumes can be analyzed."
        )

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    text = text.strip()

    if not text:
        raise ResumeAnalysisError(
            "No readable text was found in the resume."
        )

    return text[:MAX_RESUME_TEXT_LENGTH]


# --------------------------------------------------
# GEMINI
# --------------------------------------------------

def _get_client() -> "genai.Client":
    if not GEMINI_API_KEY:
        raise ResumeAnalysisError(
            "Resume analysis is not configured. "
            "GEMINI_API_KEY is missing."
        )

    return genai.Client(
        api_key=GEMINI_API_KEY
    )


def _build_prompt(
    resume_text: str,
) -> str:

    return f"""
You are an expert resume reviewer.

Analyze the following resume carefully.

RESUME:
{resume_text}

Return ONLY one JSON object.

Required structure:

{{
  "summary": "A concise professional summary of the candidate.",
  "skills": [
    "skill 1",
    "skill 2"
  ],
  "education": [
    {{
      "degree": "degree or qualification",
      "institution": "institution name",
      "year": "year or duration",
      "details": "relevant details"
    }}
  ],
  "experience": [
    {{
      "role": "job title / internship title",
      "organization": "organization name",
      "duration": "duration if available",
      "details": "main responsibilities or achievements"
    }}
  ],
  "projects": [
    {{
      "name": "project name",
      "technologies": [
        "technology 1",
        "technology 2"
      ],
      "details": "project description"
    }}
  ],
  "certifications": [
    "certification 1",
    "certification 2"
  ],
  "strengths": [
    "strength 1",
    "strength 2"
  ],
  "areas_for_improvement": [
    "improvement 1",
    "improvement 2"
  ],
  "overall_assessment": "A concise overall assessment of the resume.",
  "suggested_roles": [
    "role 1",
    "role 2"
  ]
}}

Rules:
- Use only information supported by the resume.
- Do not invent projects, skills, companies, degrees, or experience.
- Keep the analysis concise and useful.
- If a section is missing, return an empty array.
- Do not return markdown.
- Do not return code fences.
- Return ONLY valid JSON.
"""


# --------------------------------------------------
# JSON PARSING
# --------------------------------------------------

def _extract_json(
    raw_text: str,
) -> dict:

    text = raw_text.strip()

    text = re.sub(
        r"^```(?:json)?\s*",
        "",
        text,
        flags=re.IGNORECASE,
    )

    text = re.sub(
        r"\s*```$",
        "",
        text,
    )

    text = text.strip()

    try:
        data = json.loads(text)

        if isinstance(data, dict):
            return data

    except json.JSONDecodeError:
        pass

    match = re.search(
        r"\{.*\}",
        text,
        re.DOTALL,
    )

    if match:
        try:
            data = json.loads(
                match.group(0)
            )

            if isinstance(data, dict):
                return data

        except json.JSONDecodeError:
            pass

    raise ResumeAnalysisError(
        "Could not parse the resume analysis response."
    )


# --------------------------------------------------
# VALIDATION
# --------------------------------------------------

def _string_list(
    value,
) -> list[str]:

    if not isinstance(
        value,
        list,
    ):
        return []

    return [
        str(item).strip()
        for item in value
        if str(item).strip()
    ]


def _validate_analysis(
    data: dict,
) -> dict:

    if not isinstance(data, dict):
        raise ResumeAnalysisError(
            "Resume analysis returned invalid data."
        )

    education = []

    for item in data.get(
        "education",
        [],
    ):
        if not isinstance(item, dict):
            continue

        education.append(
            {
                "degree": str(
                    item.get(
                        "degree",
                        "",
                    )
                ).strip(),

                "institution": str(
                    item.get(
                        "institution",
                        "",
                    )
                ).strip(),

                "year": str(
                    item.get(
                        "year",
                        "",
                    )
                ).strip(),

                "details": str(
                    item.get(
                        "details",
                        "",
                    )
                ).strip(),
            }
        )

    experience = []

    for item in data.get(
        "experience",
        [],
    ):
        if not isinstance(item, dict):
            continue

        experience.append(
            {
                "role": str(
                    item.get(
                        "role",
                        "",
                    )
                ).strip(),

                "organization": str(
                    item.get(
                        "organization",
                        "",
                    )
                ).strip(),

                "duration": str(
                    item.get(
                        "duration",
                        "",
                    )
                ).strip(),

                "details": str(
                    item.get(
                        "details",
                        "",
                    )
                ).strip(),
            }
        )

    projects = []

    for item in data.get(
        "projects",
        [],
    ):
        if not isinstance(item, dict):
            continue

        projects.append(
            {
                "name": str(
                    item.get(
                        "name",
                        "",
                    )
                ).strip(),

                "technologies": _string_list(
                    item.get(
                        "technologies",
                        [],
                    )
                ),

                "details": str(
                    item.get(
                        "details",
                        "",
                    )
                ).strip(),
            }
        )

    return {
        "summary": str(
            data.get(
                "summary",
                "",
            )
        ).strip(),

        "skills": _string_list(
            data.get(
                "skills",
                [],
            )
        ),

        "education": education,

        "experience": experience,

        "projects": projects,

        "certifications": _string_list(
            data.get(
                "certifications",
                [],
            )
        ),

        "strengths": _string_list(
            data.get(
                "strengths",
                [],
            )
        ),

        "areas_for_improvement": _string_list(
            data.get(
                "areas_for_improvement",
                [],
            )
        ),

        "overall_assessment": str(
            data.get(
                "overall_assessment",
                "",
            )
        ).strip(),

        "suggested_roles": _string_list(
            data.get(
                "suggested_roles",
                [],
            )
        ),
    }


# --------------------------------------------------
# MAIN ANALYZER
# --------------------------------------------------

def analyze_resume(
    user_id: int,
    stored_filename: str,
) -> dict:

    path = resume_file_path(
        user_id,
        stored_filename,
    )

    resume_text = extract_resume_text(
        path
    )

    client = _get_client()

    last_error: Exception | None = None

    for attempt in range(
        1,
        MAX_ANALYSIS_ATTEMPTS + 1,
    ):

        prompt = _build_prompt(
            resume_text
        )

        try:
            response = client.models.generate_content(
                model=GEMINI_MODEL,
                contents=prompt,
                config={
                    "response_mime_type": "application/json",
                    "temperature": 0.2,
                },
            )

        except genai_errors.APIError as exc:
            logger.warning(
                "Gemini resume analysis API error "
                "attempt %s/%s: %s",
                attempt,
                MAX_ANALYSIS_ATTEMPTS,
                exc,
            )

            last_error = ResumeAnalysisError(
                "The resume analysis service is temporarily unavailable."
            )

            continue

        except Exception as exc:
            logger.exception(
                "Unexpected resume analysis error "
                "attempt %s/%s",
                attempt,
                MAX_ANALYSIS_ATTEMPTS,
            )

            last_error = ResumeAnalysisError(
                "An unexpected error occurred during resume analysis."
            )

            continue

        raw_text = getattr(
            response,
            "text",
            None,
        )

        if not raw_text:
            last_error = ResumeAnalysisError(
                "The resume analysis service returned no content."
            )

            continue

        try:
            data = _extract_json(
                raw_text
            )

            return _validate_analysis(
                data
            )

        except ResumeAnalysisError as exc:
            logger.warning(
                "Resume analysis validation failed "
                "attempt %s/%s: %s",
                attempt,
                MAX_ANALYSIS_ATTEMPTS,
                exc,
            )

            last_error = exc

            continue

    raise (
        last_error
        or ResumeAnalysisError(
            "Resume analysis failed after multiple attempts."
        )
    )